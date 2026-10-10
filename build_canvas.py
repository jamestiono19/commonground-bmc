"""CommonGround Business Model Canvas – V2 (purple / pink card design).

Outputs (same folder):
  CommonGround_Business_Model_Canvas.pptx  -> import into Canva (text, cards & pills stay editable)
  CommonGround_Business_Model_Canvas.png / .pdf -> previews
  assets/*.png -> icons + decorations (embedded in the PPTX as movable images)
All coordinates are inches on a 15 x 10 in page."""
import html, json, urllib.request
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = Path(__file__).parent
ASSETS = OUT / "assets"
ASSETS.mkdir(exist_ok=True)
W, H = 15, 10
HEAD, BODY = "Poppins", "Open Sans"

INK, PURPLE, PURPLE2, MUTED = "1E1B4B", "6D28D9", "A855F7", "5B5675"
PAGE, BORDER, PILL = "FBFAFE", "E4DAF7", "E9E0FB"
PINK = ("FCE7F3", "FDF4FA")      # gradient start, end
LAV = ("EDE7FB", "F8F5FE")
SOFT = ("F6F3FC", "FFFFFF")

# ------------------------------------------------------------------ content (simplified)
CARDS = [
    # key, x, y, w, h, bg, icon, title, question, bullets, pill
    ("kp", 0.50, 2.05, 2.68, 4.80, LAV, "handshake", "Key Partners", "Who helps us deliver value?",
     ["Ad network (e.g. Google AdSense)", "Map & place-data providers", "BINUS student organisations",
      "Campus & Festival organisers"], "Partners bring reach & data."),
    ("ka", 3.33, 2.05, 2.68, 2.33, PINK, "settings", "Key Activities", "What must we do well?",
     ["Build the 7-screen decision flow", "Filter > score > fairness > Top 3", "Verify 30–50 venues near campus",
      "Run the 10-group pilot"], None),
    ("kr", 3.33, 4.52, 2.68, 2.33, LAV, "package", "Key Resources", "What do we need?",
     ["Rules-based matching engine", "Verified venue database", "Product team",
      "Traffic from shared room links"], None),
    ("vp", 6.16, 2.05, 2.68, 4.80, SOFT, "gem", "Value Propositions", "Why do groups choose us?",
     ["Fair: hard limits first, nobody below 60", "Private: nobody sees your limits",
      "Fast: one link, no sign-up", "Advertisers: brand-safe 18–25 audience"], "Kumpulkan preferensi tanpa debat."),
    ("cr", 8.99, 2.05, 2.68, 2.33, LAV, "message-circle", "Customer Relationships", "How do we keep users?",
     ["Self-service, no sign-up", "Automated timer, Top 3 & veto", "In-app feedback & venue flags",
      "Campus student communities"], None),
    ("ch", 8.99, 4.52, 2.68, 2.33, PINK, "megaphone", "Channels", "How do we reach users?",
     ["Room links: WhatsApp, LINE, Discord", "Campus QR codes & Festival demo", "Social media",
      "Every room brings 3–7 new people"], None),
    ("cs", 11.82, 2.05, 2.68, 4.80, LAV, "users", "Customer Segments", "Who do we serve?",
     ["Users (free): BINUS students & friend groups of 4–8",
      "Payers: advertisers via ad network – games, apps, e-commerce",
      "Later: frequent hosts (CommonGround+)"], None),
    ("co", 0.50, 7.00, 6.92, 2.40, PINK, "coins", "Cost Structure", "What does it cost to run?",
     ["Development & hosting", "Data / API use", "Venue checks & updates", "Campus marketing"],
     "Lean build. Verified data."),
    ("re", 7.58, 7.00, 6.92, 2.40, LAV, "circle-dollar-sign", "Revenue Streams", "How do we make money?",
     ["Phase 1: free pilot, no ads", "Phase 2: in-app ads (CPM / CPC) on 2 screens only",
      "Phase 3: CommonGround+ for hosts, ~Rp15–25k / month",
      "Never food ads – picks stay neutral"], "Free for users. Ads first, subscription later."),
]
QUOTE = "\u201cSatu link. Semua jawab. Tiga pilihan pas.\u201d"
CS_TAG = "THE RIGHT GROUP,\nTHE RIGHT PLACE"
INFO = [("Business Name:", "CommonGround"), ("Date:", "10 Oct 2026"),
        ("Prepared By:", "CommonGround Team"), ("Version:", "Food-only MVP · V3")]

# ------------------------------------------------------------------ art (SVG -> PNG)
def lucide(name):
    url = f"https://cdn.jsdelivr.net/npm/lucide-static@0.460.0/icons/{name}.svg"
    svg = urllib.request.urlopen(url, timeout=30).read().decode()
    return svg[svg.index("<svg"):].replace('stroke="currentColor"', 'stroke="#FFFFFF"')


def icon_html(name):
    inner = lucide(name).replace('width="24"', 'width="62"').replace('height="24"', 'height="62"')
    return f"""<div style="width:120px;height:120px;border-radius:50%;display:flex;align-items:center;justify-content:center;
      background:linear-gradient(135deg,#7C3AED,#C084FC);box-shadow:inset 0 -6px 0 rgba(0,0,0,.08)">{inner}</div>"""


WAVES = "<svg width='780' height='560' viewBox='0 0 780 560'><defs><linearGradient id='g' x1='0' x2='1'>" \
        "<stop offset='0' stop-color='#C4B5FD'/><stop offset='.6' stop-color='#E879F9'/><stop offset='1' stop-color='#FB923C'/>" \
        "</linearGradient></defs>" + "".join(
            f"<path d='M0 {200 + i * 9} C 200 {60 + i * 14}, 420 {420 - i * 6}, 780 {150 + i * 12}' fill='none' "
            f"stroke='url(#g)' stroke-width='2' opacity='{0.35 + i * 0.03:.2f}'/>" for i in range(18)) + \
        "<circle cx='560' cy='120' r='12' fill='#FDBA74'/><circle cx='600' cy='70' r='7' fill='#A78BFA'/>" \
        "<path d='M170 470 q 30 -40 60 0' stroke='#7C3AED' stroke-width='5' fill='none' stroke-linecap='round'/></svg>"

BLOB = """<svg width='780' height='640' viewBox='0 0 780 640'><defs>
<linearGradient id='a' x1='0' y1='0' x2='1' y2='1'><stop offset='0' stop-color='#FBCFE8'/><stop offset='1' stop-color='#FDBA74'/></linearGradient>
<linearGradient id='b' x1='0' y1='0' x2='1' y2='0'><stop offset='0' stop-color='#C4B5FD'/><stop offset='1' stop-color='#7C3AED'/></linearGradient>
<clipPath id='k'><rect width='780' height='640' rx='52'/></clipPath></defs><g clip-path='url(#k)'>
<path d='M0 330 C 160 250 300 420 470 300 S 700 140 780 170 L780 640 L0 640 Z' fill='url(#b)' opacity='.85'/>
<path d='M0 420 C 180 340 330 520 520 400 S 720 300 780 330 L780 640 L0 640 Z' fill='url(#a)' opacity='.95'/></g>
""" + "".join(f"<circle cx='{90 + c * 26}' cy='{90 + r * 26}' r='5' fill='#A78BFA'/>" for r in range(3) for c in range(3)) + \
       "<circle cx='640' cy='60' r='16' fill='#FDBA74'/><circle cx='560' cy='130' r='10' fill='#7C3AED'/></svg>"

SWOOSH = """<svg width='900' height='330' viewBox='0 0 900 330'><defs>
<linearGradient id='s' x1='0' y1='0' x2='1' y2='0'><stop offset='0' stop-color='#FBCFE8' stop-opacity='0'/>
<stop offset='.5' stop-color='#F9A8D4'/><stop offset='1' stop-color='#8B5CF6'/></linearGradient>
<clipPath id='k'><rect width='900' height='330' rx='48'/></clipPath></defs><g clip-path='url(#k)'>
<path d='M120 330 C 420 300 640 200 900 30 L900 330 Z' fill='url(#s)'/>
<path d='M520 330 C 700 300 820 220 900 150 L900 330 Z' fill='#7C3AED' opacity='.75'/></g></svg>"""


def star(cx, cy, r, fill):
    import math
    pts = []
    for i in range(10):
        a = math.pi / 2 + i * math.pi / 5
        rr = r if i % 2 == 0 else r * 0.45
        pts.append(f"{cx + rr * math.cos(a):.1f},{cy - rr * math.sin(a):.1f}")
    return f"<polygon points='{' '.join(pts)}' fill='{fill}'/>"


def person(cx, color, face="#FFF7ED"):
    return (f"<rect x='{cx - 38}' y='150' width='76' height='90' rx='36' fill='{color}'/>"
            f"<circle cx='{cx}' cy='112' r='30' fill='{face}' stroke='{color}' stroke-width='5'/>"
            f"<circle cx='{cx - 10}' cy='110' r='3.5' fill='#1E1B4B'/><circle cx='{cx + 10}' cy='110' r='3.5' fill='#1E1B4B'/>"
            f"<path d='M{cx - 10} 124 q10 9 20 0' stroke='#1E1B4B' stroke-width='3' fill='none' stroke-linecap='round'/>")


PEOPLE = f"""<svg width='740' height='380' viewBox='0 0 740 380'><defs>
<linearGradient id='p' x1='0' y1='0' x2='0' y2='1'><stop offset='0' stop-color='#F3E8FF'/><stop offset='1' stop-color='#FCE7F3'/></linearGradient></defs>
<rect width='740' height='380' rx='34' fill='url(#p)'/>
{person(200, '#A78BFA')}{person(370, '#F472B6')}{person(540, '#FBBF24')}
<ellipse cx='370' cy='262' rx='290' ry='62' fill='#FFFFFF' stroke='#E9D5FF' stroke-width='4'/>
<circle cx='260' cy='262' r='14' fill='#C4B5FD'/><circle cx='370' cy='268' r='18' fill='#FDBA74' stroke='#F97316' stroke-width='4'/>
<circle cx='480' cy='262' r='14' fill='#F9A8D4'/>
<rect x='300' y='22' width='140' height='46' rx='23' fill='#7C3AED'/>
<text x='345' y='54' font-family='Poppins' font-weight='700' font-size='24' fill='#fff' text-anchor='middle'>Top 3</text>
{star(408, 45, 13, '#FDE047')}</svg>"""

ART = {"waves": WAVES, "blob": BLOB, "swoosh": SWOOSH, "people": PEOPLE}
ICONS = {c[6] for c in CARDS}


def render_assets(pg):
    for name in ICONS:
        pg.set_content(f"<body style='margin:0;background:transparent'><div id='e' style='display:inline-block'>{icon_html(name)}</div></body>")
        pg.locator("#e").screenshot(path=str(ASSETS / f"icon-{name}.png"), omit_background=True)
    for name, svg in ART.items():
        pg.set_content("<link href='https://fonts.googleapis.com/css2?family=Poppins:wght@700&display=block' rel='stylesheet'>"
                       f"<body style='margin:0;background:transparent'><div id='e' style='display:inline-block;line-height:0'>{svg}</div></body>")
        pg.wait_for_load_state("networkidle"); pg.evaluate("document.fonts.ready")
        pg.locator("#e").screenshot(path=str(ASSETS / f"{name}.png"), omit_background=True)


# ------------------------------------------------------------------ layout helpers (shared)
def layout():
    """Return list of element dicts used by both renderers."""
    E = []
    add = lambda **k: E.append(k)
    # header
    add(t="text", x=0.50, y=0.38, w=4, h=0.42, runs=[("CommonGround", PURPLE)], font=HEAD, size=19, bold=True)
    add(t="text", x=0.50, y=0.76, w=8, h=0.75, runs=[("Business Model ", PURPLE), ("Canvas", INK)], font=HEAD, size=34, bold=True)
    add(t="text", x=0.52, y=1.50, w=8, h=0.32, runs=[("A one-page view of how CommonGround creates, delivers and captures value.", MUTED)],
        font=BODY, size=10.5)
    # info box
    add(t="card", x=9.30, y=0.38, w=5.20, h=1.42, bg=("F1EBFD", "FBF8FF"), line="C4B5FD", lw=1.25)
    for i, (lab, val) in enumerate(INFO):
        col, row = i % 2, i // 2
        bx = 9.50 + col * 2.75
        lw_ = 0.95 if col == 0 else 0.48
        fw = 2.45 if col == 0 else 1.92
        by = 0.62 + row * 0.55
        add(t="text", x=bx, y=by, w=lw_ + 0.1, h=0.28, runs=[(lab, INK)], font=BODY, size=8, bold=True)
        add(t="text", x=bx + lw_, y=by - 0.02, w=fw - lw_, h=0.28, runs=[(val, PURPLE)], font=BODY, size=8.5, bold=True)
        add(t="line", x=bx + lw_, y=by + 0.27, w=fw - lw_, color="A78BFA")
    for (key, x, y, w, h, bg, icon, title, q, bullets, pill) in CARDS:
        wide = key in ("co", "re")
        add(t="card", x=x, y=y, w=w, h=h, bg=bg, line=BORDER, lw=1)
        # decorations first (under text)
        if key == "kp":
            add(t="img", src="waves", x=x + 0.05, y=y + 2.25, w=w - 0.1, h=(w - 0.1) * 560 / 780)
        if key == "cs":
            add(t="img", src="blob", x=x + 0.01, y=y + h - 0.01 - (w - 0.02) * 640 / 780, w=w - 0.02, h=(w - 0.02) * 640 / 780)
        if wide:
            add(t="img", src="swoosh", x=x + w - 3.4, y=y + h - 0.01 - 3.39 * 330 / 900, w=3.39, h=3.39 * 330 / 900)
        add(t="img", src=f"icon-{icon}", x=x + 0.16, y=y + 0.16, w=0.50, h=0.50)
        add(t="text", x=x + 0.74, y=y + 0.16, w=w - 0.85, h=0.50, runs=[(title, INK)], font=HEAD,
            size=12 if len(title) > 18 else 13, bold=True, anchor="m")
        qy = y + 0.78
        if wide:
            add(t="text", x=x + 0.74, y=y + 0.70, w=w - 0.9, h=0.3, runs=[(q, INK)], font=BODY, size=9.5, bold=True)
            add(t="bullets", x=x + 0.74, y=y + 0.98, w=w * 0.62, h=0.92, items=bullets, size=8.8)
        else:
            add(t="text", x=x + 0.18, y=qy, w=w - 0.34, h=0.3, runs=[(q, INK)], font=BODY, size=9.5, bold=True)
            bh = 1.2 if h < 3 else (1.6 if key != "cs" else 1.75)
            add(t="bullets", x=x + 0.18, y=qy + 0.32, w=w - 0.30, h=bh, items=bullets, size=8.6)
        if key == "vp":
            add(t="img", src="people", x=x + 0.18, y=y + 2.62, w=w - 0.36, h=(w - 0.36) * 380 / 740)
            add(t="text", x=x + 0.12, y=y + 3.84, w=w - 0.24, h=0.52, runs=[(QUOTE, INK)], font=HEAD, size=10.5,
                bold=True, align="c")
        if key == "cs":
            add(t="text", x=x + 0.2, y=y + h - 0.70, w=w - 0.4, h=0.5, runs=[(CS_TAG, INK)], font=HEAD, size=8,
                bold=True, align="c", spacing=1.5)
        if pill:
            pw = min(w - 0.36, 0.07 * len(pill) + 0.35)
            px_ = x + 0.18 if not wide else x + 0.74
            add(t="pill", x=px_, y=y + h - (0.42 if wide else 0.48), w=pw, h=0.30, text=pill)
    add(t="text", x=0.50, y=9.55, w=14, h=0.3, runs=[(
        "Source: CommonGround Team Brief, Food-only MVP · V3 – all figures are hypotheses to validate. "
        "Based on the Business Model Canvas by Strategyzer (CC BY-SA 3.0).", "8A85A3")], font=BODY, size=7)
    return E


# ------------------------------------------------------------------ PPTX
rgb = RGBColor.from_string
AL = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER}
AN = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE}


def tf_setup(tf, anchor="t", m=0.02):
    tf.word_wrap = True; tf.auto_size = None; tf.vertical_anchor = AN[anchor]
    for s in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, s, Inches(m))


def build_pptx(E):
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bgc = sl.background.fill; bgc.solid(); bgc.fore_color.rgb = rgb(PAGE)
    for e in E:
        X, Y, Wd = Inches(e["x"]), Inches(e["y"]), Inches(e["w"])
        if e["t"] == "card":
            s = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X, Y, Wd, Inches(e["h"]))
            s.adjustments[0] = 0.18 / min(e["w"], e["h"])
            s.fill.gradient(); s.fill.gradient_angle = 45
            s.fill.gradient_stops[0].color.rgb = rgb(e["bg"][0]); s.fill.gradient_stops[1].color.rgb = rgb(e["bg"][1])
            s.line.color.rgb = rgb(e["line"]); s.line.width = Pt(e["lw"]); s.shadow.inherit = False
        elif e["t"] == "img":
            sl.shapes.add_picture(str(ASSETS / f"{e['src']}.png"), X, Y, Wd, Inches(e["h"]))
        elif e["t"] == "line":
            c = sl.shapes.add_connector(1, X, Y, X + Wd, Y); c.line.color.rgb = rgb(e["color"]); c.line.width = Pt(0.75)
        elif e["t"] == "pill":
            s = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X, Y, Wd, Inches(e["h"]))
            s.adjustments[0] = 0.5; s.fill.solid(); s.fill.fore_color.rgb = rgb(PILL)
            s.line.fill.background(); s.shadow.inherit = False
            tf = s.text_frame; tf_setup(tf, "m", 0.05)
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            r = p.add_run(); r.text = e["text"]; f = r.font
            f.name = BODY; f.size = Pt(7.8); f.bold = True; f.color.rgb = rgb("4C1D95")
        elif e["t"] == "text":
            tb = sl.shapes.add_textbox(X, Y, Wd, Inches(e["h"])); tf = tb.text_frame
            tf_setup(tf, e.get("anchor", "t"))
            lines = [[]]
            for txt, col in e["runs"]:
                parts = txt.split("\n")
                for i, ptxt in enumerate(parts):
                    if i: lines.append([])
                    lines[-1].append((ptxt, col))
            for i, ln in enumerate(lines):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = AL[e.get("align", "l")]
                for txt, col in ln:
                    r = p.add_run(); r.text = txt; f = r.font
                    f.name = e["font"]; f.size = Pt(e["size"]); f.bold = e.get("bold", False); f.color.rgb = rgb(col)
        elif e["t"] == "bullets":
            tb = sl.shapes.add_textbox(X, Y, Wd, Inches(e["h"])); tf = tb.text_frame; tf_setup(tf)
            for i, it in enumerate(e["items"]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(3)
                r = p.add_run(); r.text = "•  " + it; f = r.font
                f.name = BODY; f.size = Pt(e["size"]); f.color.rgb = rgb("2E2A47")
    path = OUT / "CommonGround_Business_Model_Canvas.pptx"; prs.save(path); return path


# ------------------------------------------------------------------ HTML preview
def build_html(E):
    px = lambda v: f"{v * 96:.1f}px"
    out = []
    for e in E:
        pos = f"left:{px(e['x'])};top:{px(e['y'])};width:{px(e['w'])};"
        if e["t"] == "card":
            out.append(f"<div style='{pos}height:{px(e['h'])};border-radius:{px(0.18)};border:{e['lw']}pt solid #{e['line']};"
                       f"background:linear-gradient(135deg,#{e['bg'][0]},#{e['bg'][1]});box-sizing:border-box'></div>")
        elif e["t"] == "img":
            out.append(f"<img src='assets/{e['src']}.png' style='position:absolute;{pos}height:{px(e['h'])}'>")
        elif e["t"] == "line":
            out.append(f"<div style='{pos}height:0;border-top:0.75pt solid #{e['color']}'></div>")
        elif e["t"] == "pill":
            out.append(f"<div class='chk' style='{pos}height:{px(e['h'])};border-radius:99px;background:#{PILL};font:700 7.8pt {BODY};"
                       f"color:#4C1D95;display:flex;align-items:center;justify-content:center;white-space:nowrap;overflow:hidden'>{html.escape(e['text'])}</div>")
        elif e["t"] == "text":
            inner = "".join(f"<span style='color:#{c}'>{html.escape(t).replace(chr(10), '<br>')}</span>" for t, c in e["runs"])
            st = (f"{pos}height:{px(e['h'])};font-family:'{e['font']}';font-size:{e['size']}pt;font-weight:{700 if e.get('bold') else 400};"
                  f"text-align:{'center' if e.get('align') == 'c' else 'left'};padding:{px(0.02)};"
                  + ("display:flex;align-items:center;" if e.get("anchor") == "m" else ""))
            out.append(f"<div class='chk' style=\"{st}\"><div>{inner}</div></div>")
        elif e["t"] == "bullets":
            inner = "".join(f"<p>•&nbsp;&nbsp;{html.escape(i)}</p>" for i in e["items"])
            out.append(f"<div class='chk' style=\"{pos}height:{px(e['h'])};font:{e['size']}pt '{BODY}';color:#2E2A47;padding:{px(0.02)}\">{inner}</div>")
    doc = f"""<!doctype html><html><head><meta charset='utf-8'>
<link href='https://fonts.googleapis.com/css2?family=Poppins:wght@400;700&family=Open+Sans:wght@400;700&display=block' rel='stylesheet'>
<style>@page{{size:{W}in {H}in;margin:0}} body{{margin:0}}
#c{{position:relative;width:{px(W)};height:{px(H)};background:#{PAGE};overflow:hidden}}
#c>div{{position:absolute;box-sizing:border-box;line-height:1.22;overflow:hidden}} p{{margin:0 0 3pt 0}}</style></head>
<body><div id='c'>{''.join(out)}</div></body></html>"""
    p = OUT / "preview.html"; p.write_text(doc, encoding="utf-8"); return p


if __name__ == "__main__":
    from playwright.sync_api import sync_playwright
    E = layout()
    with sync_playwright() as pw:
        b = pw.chromium.launch(); pg = b.new_page(device_scale_factor=3)
        render_assets(pg)
        print("pptx:", build_pptx(E))
        hp = build_html(E)
        pg = b.new_page(viewport={"width": W * 96, "height": H * 96}, device_scale_factor=2)
        pg.goto(hp.resolve().as_uri(), wait_until="networkidle"); pg.evaluate("document.fonts.ready")
        over = pg.evaluate("[...document.querySelectorAll('.chk')].filter(e=>e.scrollHeight>e.clientHeight+1||e.scrollWidth>e.clientWidth+1).map(e=>e.innerText.slice(0,40))")
        pg.screenshot(path=str(OUT / "CommonGround_Business_Model_Canvas.png"))
        pg.pdf(path=str(OUT / "CommonGround_Business_Model_Canvas.pdf"), width=f"{W}in", height=f"{H}in", print_background=True)
        b.close()
    print("overflow:", json.dumps(over) if over else "none")
